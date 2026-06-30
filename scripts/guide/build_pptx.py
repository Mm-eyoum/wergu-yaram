# -*- coding: utf-8 -*-
"""
Construit la présentation PowerPoint « Guide d'utilisateur Wergu Yaram ».
Moteur de rendu brandé (couleurs, typo Inter, logo) + gabarits de slides.
Le contenu provient de content.py.

Usage : python3 scripts/guide/build_pptx.py
Sortie : docs/guide/Wergu_Yaram_Guide_Utilisateur.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

import content as C

try:
    from PIL import Image
except Exception:
    Image = None

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SHOTS = os.path.join(ROOT, "docs", "guide", "screenshots")
LOGO = os.path.join(ROOT, "public", "logo.png")
OUT = os.path.join(ROOT, "docs", "guide", "Wergu_Yaram_Guide_Utilisateur.pptx")

# ------------------------------------------------------------------ BRANDING
GREEN = RGBColor(0x00, 0x7A, 0x5E)
GREEN_DARK = RGBColor(0x00, 0x6B, 0x52)
GREEN_LT = RGBColor(0x00, 0xA8, 0x78)
TEAL = RGBColor(0x00, 0xB8, 0x94)
NAVY = RGBColor(0x0B, 0x1F, 0x49)
MINT = RGBColor(0xEE, 0xFD, 0xF8)
SOFT = RGBColor(0xF7, 0xFB, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TXT = RGBColor(0x0B, 0x1F, 0x49)
TXT2 = RGBColor(0x66, 0x70, 0x85)
BORDER = RGBColor(0xE6, 0xEE, 0xF2)
DANGER = RGBColor(0xC8, 0x1E, 0x1E)
DANGER_BG = RGBColor(0xFD, 0xEC, 0xEC)
WARN = RGBColor(0xF6, 0xB4, 0x4B)

FONT = "Inter"

# Couleur d'accent par rôle
ROLE = {
    "brand": GREEN,
    "public": GREEN,
    "pro": TEAL,
    "partner": NAVY,
    "admin": RGBColor(0x4F, 0x39, 0xC7),  # violet-indigo pour distinguer l'admin
}
ROLE_LABEL = {
    "brand": "WERGU YARAM",
    "public": "GRAND PUBLIC & PATIENTS",
    "pro": "PROFESSIONNELS DE SANTÉ",
    "partner": "PARTENAIRES",
    "admin": "ADMINISTRATEURS",
}

SW, SH = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]

_pageno = {"n": 0}


# ------------------------------------------------------------------ HELPERS
def slide():
    return prs.slides.add_slide(BLANK)


def _hex(c):
    return "%02X%02X%02X" % (c[0], c[1], c[2])


def rect(s, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE,
         radius=None, shadow=False):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if shadow:
        soft_shadow(sp)
    return sp


def soft_shadow(shape, blur=110000, dist=40000, alpha=86):
    """Ombre douce style 'card' (navy très transparent)."""
    try:
        spPr = shape._element.spPr
        for el in spPr.findall(qn("a:effectLst")):
            spPr.remove(el)
        a = int((100 - alpha) * 1000)  # alpha faible = très transparent
        xml = (
            '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<a:outerShdw blurRad="{blur}" dist="{dist}" dir="5400000" rotWithShape="0">'
            f'<a:srgbClr val="0B1F49"><a:alpha val="{a}"/></a:srgbClr>'
            "</a:outerShdw></a:effectLst>"
        )
        spPr.append(parse_xml(xml))
    except Exception:
        pass


def gradient(shape, c1, c2, angle=45):
    """Remplit une forme avec un dégradé linéaire c1->c2."""
    try:
        spPr = shape._element.spPr
        for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
                    "a:pattFill", "a:grpFill"):
            for el in spPr.findall(qn(tag)):
                spPr.remove(el)
        ang = int(angle * 60000)
        xml = (
            '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" rotWithShape="1">'
            "<a:gsLst>"
            f'<a:gs pos="0"><a:srgbClr val="{_hex(c1)}"/></a:gs>'
            f'<a:gs pos="100000"><a:srgbClr val="{_hex(c2)}"/></a:gs>'
            "</a:gsLst>"
            f'<a:lin ang="{ang}" scaled="1"/></a:gradFill>'
        )
        geom = spPr.find(qn("a:prstGeom"))
        if geom is None:
            geom = spPr.find(qn("a:custGeom"))
        grad = parse_xml(xml)
        if geom is not None:
            geom.addnext(grad)
        else:
            spPr.append(grad)
    except Exception:
        shape.fill.solid()
        shape.fill.fore_color.rgb = c1


def text(s, x, y, w, h, runs, size=14, color=TXT, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, italic=False, spacing=1.0, wrap=True):
    """runs : str OU liste de paragraphes (str) OU liste de (str, dict)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [runs]
    for i, item in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        opts = {}
        if isinstance(item, tuple):
            item, opts = item
        r = p.add_run()
        r.text = item
        f = r.font
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.italic = opts.get("italic", italic)
        f.name = opts.get("font", font)
        f.color.rgb = opts.get("color", color)
    return tb


def fit(box_w, box_h, img_w, img_h):
    r = min(box_w / img_w, box_h / img_h)
    return img_w * r, img_h * r


def place_image(s, path, bx, by, bw, bh, framed=True):
    """Place une capture dans un cadre blanc arrondi avec ombre, centrée."""
    if framed:
        card = rect(s, bx, by, bw, bh, fill=WHITE, line=BORDER, line_w=1.0,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.045, shadow=True)
    pad = 0.12
    iw_box, ih_box = bw - 2 * pad, bh - 2 * pad
    if path and os.path.exists(path) and Image is not None:
        try:
            with Image.open(path) as im:
                w0, h0 = im.size
        except Exception:
            w0, h0 = 1440, 900
        w, h = fit(iw_box, ih_box, w0, h0)
        ix = bx + (bw - w) / 2
        iy = by + (bh - h) / 2
        s.shapes.add_picture(path, Inches(ix), Inches(iy), Inches(w), Inches(h))
    else:
        # placeholder
        rect(s, bx + pad, by + pad, iw_box, ih_box, fill=MINT,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
        text(s, bx, by, bw, bh, "Aperçu de l'écran", size=14, color=TXT2,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def logo_image(s, x, y, h):
    if os.path.exists(LOGO) and Image is not None:
        try:
            with Image.open(LOGO) as im:
                w0, h0 = im.size
            w = h * (w0 / h0)
            s.shapes.add_picture(LOGO, Inches(x), Inches(y), Inches(w), Inches(h))
            return w
        except Exception:
            pass
    return 0


def footer(s, accent=GREEN):
    _pageno["n"] += 1
    rect(s, 0.6, 7.02, 12.13, 0.012, fill=BORDER)
    text(s, 0.6, 7.08, 8, 0.3, "Wergu Yaram — Portail santé du Sénégal",
         size=9, color=TXT2)
    text(s, 10.0, 7.08, 2.73, 0.3, str(_pageno["n"]), size=9, color=accent,
         bold=True, align=PP_ALIGN.RIGHT)


def pill(s, x, y, d, label, fill, txtcolor=WHITE, size=14):
    rect(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    text(s, x, y - 0.01, d, d, label, size=size, color=txtcolor, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def chip(s, x, y, label, accent):
    w = 0.18 + 0.092 * len(label)
    rect(s, x, y, w, 0.34, fill=MINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(s, x, y - 0.005, w, 0.34, label.upper(), size=9, color=accent, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


# ------------------------------------------------------------------ SLIDES
def s_cover(d):
    s = slide()
    bg = rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1)
    gradient(bg, TEAL, GREEN, angle=120)
    # motifs décoratifs (cercles translucides)
    for (cx, cy, dd, c) in [(11.2, -1.5, 5.5, GREEN_LT), (12.3, 5.4, 3.2, TEAL),
                            (-1.2, 5.8, 3.6, GREEN_LT)]:
        o = rect(s, cx, cy, dd, dd, fill=c, shape=MSO_SHAPE.OVAL)
        o.fill.fore_color.rgb = c
        try:
            o.fill.transparency = 0.7
        except Exception:
            pass
    # carte logo blanche
    lw = 3.4
    rect(s, 0.9, 0.85, lw, 1.25, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         radius=0.18, shadow=True)
    logo_image(s, 1.12, 1.12, 0.72)
    # titres
    text(s, 0.9, 2.95, 11, 0.5, ("GUIDE D'UTILISATEUR"),
         size=16, color=WHITE, bold=True)
    text(s, 0.88, 3.4, 11.5, 1.6, d.get("title", "Guide d'utilisateur"),
         size=58, color=WHITE, bold=True, spacing=1.0)
    text(s, 0.9, 5.0, 11, 0.6, d.get("tagline", ""), size=24, color=MINT, bold=True)
    text(s, 0.9, 5.6, 10.5, 0.8, d.get("subtitle", ""), size=15, color=WHITE)
    text(s, 0.9, 6.75, 11.5, 0.4, d.get("footer", ""), size=12, color=MINT)


def s_closing(d):
    s = slide()
    bg = rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1)
    gradient(bg, GREEN, TEAL, angle=120)
    logo_card = rect(s, SW / 2 - 1.7, 1.5, 3.4, 1.25, fill=WHITE,
                     shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18, shadow=True)
    logo_image(s, SW / 2 - 1.48, 1.77, 0.72)
    text(s, 0.5, 3.5, SW - 1, 1.0, d.get("title", "Merci"), size=42, color=WHITE,
         bold=True, align=PP_ALIGN.CENTER)
    text(s, 1.5, 4.7, SW - 3, 0.8, d.get("subtitle", ""), size=18, color=MINT,
         align=PP_ALIGN.CENTER)
    text(s, 0.5, 6.6, SW - 1, 0.4, d.get("footer", ""), size=12, color=WHITE,
         align=PP_ALIGN.CENTER)


def s_toc(d):
    s = slide()
    rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1, fill=SOFT)
    rect(s, 0, 0, 0.28, SH, fill=GREEN)
    text(s, 0.7, 0.55, 11, 0.8, d["title"], size=32, color=NAVY, bold=True)
    rect(s, 0.72, 1.35, 0.9, 0.07, fill=TEAL)
    y = 1.75
    rh = 0.72
    for (num, title, sub, role) in d["items"]:
        accent = ROLE.get(role, GREEN)
        pill(s, 0.72, y, 0.5, num, accent, size=16)
        text(s, 1.45, y - 0.04, 10.5, 0.4, title, size=17, color=NAVY, bold=True)
        text(s, 1.45, y + 0.34, 10.8, 0.34, sub, size=11, color=TXT2)
        y += rh
    footer(s)


def s_section(d):
    s = slide()
    rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1, fill=NAVY)
    accent = ROLE.get(d.get("role", "brand"), GREEN)
    band = rect(s, -0.05, -0.05, 4.6, SH + 0.1)
    gradient(band, accent, NAVY, angle=120)
    # gros numéro
    text(s, 0.7, 1.2, 3.5, 3.0, d["num"], size=200, color=WHITE, bold=True)
    try:
        # léger voile
        pass
    except Exception:
        pass
    chip_w = chip(s, 5.1, 2.45, ROLE_LABEL.get(d.get("role", "brand"), ""), accent) \
        if d.get("role") in ROLE_LABEL else 0
    text(s, 5.05, 2.95, 7.7, 1.6, d["title"], size=40, color=WHITE, bold=True,
         spacing=1.0)
    rect(s, 5.12, 4.35, 1.0, 0.08, fill=accent)
    text(s, 5.05, 4.6, 7.6, 1.0, d.get("subtitle", ""), size=16, color=MINT)
    logo_image(s, 5.05, 6.55, 0.42)


def s_feature(d):
    s = slide()
    rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1, fill=WHITE)
    accent = ROLE.get(d.get("role", "brand"), GREEN)
    rect(s, 0, 0, 0.28, SH, fill=accent)
    # bandeau rôle + titre
    chip(s, 0.7, 0.5, ROLE_LABEL.get(d.get("role", "brand"), ""), accent)
    text(s, 0.7, 0.95, 6.1, 0.9, d["title"], size=27, color=NAVY, bold=True,
         spacing=1.0)
    text(s, 0.7, 1.95, 5.95, 1.1, d.get("intro", ""), size=12.5, color=TXT2,
         spacing=1.05)

    # étapes
    y = 3.05
    steps = d.get("steps", [])
    rh = 0.78 if len(steps) <= 4 else 0.66
    for i, st in enumerate(steps, 1):
        pill(s, 0.72, y, 0.42, str(i), accent, size=13)
        text(s, 1.32, y - 0.06, 5.35, rh, st, size=12, color=TXT, spacing=1.02,
             anchor=MSO_ANCHOR.MIDDLE)
        y += rh

    # encadré astuce / important
    box_y = 6.05
    if d.get("tip"):
        rect(s, 0.7, box_y, 5.95, 0.82, fill=MINT, line=None,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(s, 0.92, box_y + 0.1, 5.6, 0.25, "ASTUCE", size=9.5, color=GREEN, bold=True)
        text(s, 0.92, box_y + 0.33, 5.55, 0.45, d["tip"], size=10.5, color=TXT,
             spacing=1.0)
    elif d.get("important"):
        rect(s, 0.7, box_y, 5.95, 0.82, fill=DANGER_BG, line=None,
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        text(s, 0.92, box_y + 0.1, 5.6, 0.25, "IMPORTANT", size=9.5, color=DANGER, bold=True)
        text(s, 0.92, box_y + 0.33, 5.55, 0.45, d["important"], size=10.5, color=TXT,
             spacing=1.0)

    # capture à droite
    img = os.path.join(SHOTS, d["img"]) if d.get("img") else None
    place_image(s, img, 6.95, 1.0, 5.78, 5.65)
    footer(s, accent)


def s_bullets(d):
    s = slide()
    rect(s, -0.05, -0.05, SW + 0.1, SH + 0.1, fill=SOFT)
    accent = ROLE.get(d.get("role", "brand"), GREEN)
    rect(s, 0, 0, 0.28, SH, fill=accent)
    chip(s, 0.7, 0.5, ROLE_LABEL.get(d.get("role", "brand"), ""), accent)
    text(s, 0.7, 0.95, 12, 0.7, d["title"], size=27, color=NAVY, bold=True)
    text(s, 0.7, 1.7, 11.8, 0.8, d.get("intro", ""), size=13, color=TXT2, spacing=1.05)

    cols = d.get("columns", [])
    n = len(cols)
    # grille : 2 colonnes si <=4, sinon 3
    per_row = 2 if n <= 4 else 3
    gap = 0.3
    total_w = 12.0
    cw = (total_w - gap * (per_row - 1)) / per_row
    ch = 1.95 if n <= 4 else 1.7
    x0, y0 = 0.7, 2.75
    for i, (h, body) in enumerate(cols):
        r = i // per_row
        cc = i % per_row
        x = x0 + cc * (cw + gap)
        y = y0 + r * (ch + 0.3)
        card = rect(s, x, y, cw, ch, fill=WHITE, line=BORDER, line_w=1.0,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07, shadow=True)
        rect(s, x + 0.32, y + 0.32, 0.5, 0.07, fill=accent)
        text(s, x + 0.32, y + 0.5, cw - 0.6, 0.55, h, size=15, color=NAVY, bold=True,
             spacing=1.0)
        text(s, x + 0.32, y + 1.05, cw - 0.6, ch - 1.2, body, size=11, color=TXT2,
             spacing=1.05)
    footer(s, accent)


# ------------------------------------------------------------------ BUILD
RENDER = {
    "cover": s_cover,
    "toc": s_toc,
    "section": s_section,
    "feature": s_feature,
    "bullets": s_bullets,
    "closing": s_closing,
}

missing = []
for d in C.SLIDES:
    if d.get("img"):
        p = os.path.join(SHOTS, d["img"])
        if not os.path.exists(p):
            missing.append(d["img"])
    RENDER[d["type"]](d)

prs.save(OUT)
print(f"✓ Présentation générée : {OUT}")
print(f"  {len(C.SLIDES)} slides · {len(prs.slides._sldIdLst)} slides écrites")
if missing:
    print(f"  ⚠︎ Captures manquantes (placeholder utilisé) : {missing}")
